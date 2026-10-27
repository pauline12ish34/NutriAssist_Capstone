"""Gemini drafts: reply parsing, reviewer checks, and approved drafts reaching the training split only."""

import json

import pytest
from openpyxl import load_workbook

from nutriassist.qa_data.build_qa_dataset import augmented_examples
from nutriassist.qa_data.gemini_drafts import (
    api_key, apply_decisions, build_prompt, checks, export_sheet, load_drafts, parse_response, source_pairs,
)

MASTER = "data/QA_data/processed/qa_master.jsonl"


def master():
    with open(MASTER, encoding="utf-8") as handle:
        return {row["id"]: row for row in map(json.loads, handle)}


@pytest.fixture
def pair():
    return next(p for p in source_pairs(MASTER) if any(c.isdigit() for c in p["answer"]))


@pytest.fixture
def drafts(tmp_path, pair):
    path = tmp_path / "drafts.jsonl"
    rows = [{"id": pair["id"], "model": "test", "questions": ["Ikibazo gishya?", "Ikibazo  gishya?", pair["question"]],
             "answers": ["Igisubizo gishya."]}]
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows), encoding="utf-8")
    return path


def test_only_reviewed_kinyarwanda_training_pairs_are_drafted():
    pairs = source_pairs(MASTER)
    assert pairs and all(p["lang"] == "rw" and p["split"] == "train" for p in pairs)
    assert sum(bool(p["english_question"]) for p in pairs) > len(pairs) / 2   # a few pairs exist only in Kinyarwanda


def test_prompt_carries_the_pairs_and_the_glossary(pair):
    prompt = build_prompt([pair])
    assert pair["id"] in prompt and pair["answer"] in prompt
    assert "breast milk = amashereka" in prompt and "gusya" in prompt


def test_parse_response_keeps_known_ids_only():
    reply = '```json\n[{"id": "rw-a", "questions": ["Q1 ", ""], "answer": "A1"}, {"id": "rw-x", "questions": ["Q"]}]\n```'
    assert parse_response(reply, {"rw-a"}) == {"rw-a": {"questions": ["Q1"], "answers": ["A1"]}}


def test_checks_flag_changed_numbers_and_gusekura(pair):
    assert checks("answer", pair["answer"], pair) == ""
    assert checks("answer", pair["answer"] + " 999", pair) == "numbers differ"
    assert checks("question", "Nshobora gusekura ibishyimbo?", pair) == "uses gusekura"


def test_drafts_drop_duplicates_and_existing_questions(drafts, pair):
    assert load_drafts(drafts, MASTER) == [(pair["id"], "question", "Ikibazo gishya?"),
                                           (pair["id"], "answer", "Igisubizo gishya.")]


def test_api_key_is_read_from_env_file(tmp_path, monkeypatch):
    for name in ("GEMINI_API_KEY", "GOOGLE_API_KEY", "GOOGLE_APIKEY"):
        monkeypatch.delenv(name, raising=False)
    env = tmp_path / ".env"
    env.write_text('other=1\ngoogle_apikey="abc123"\n', encoding="utf-8")
    assert api_key(env) == "abc123"
    with pytest.raises(RuntimeError, match="No Gemini API key"):
        api_key(tmp_path / "missing.env")


def test_approved_drafts_become_training_examples(drafts, tmp_path, pair):
    sheet, output = tmp_path / "review.xlsx", tmp_path / "augmented.jsonl"
    export_sheet(drafts, sheet, MASTER)
    book = load_workbook(sheet)
    rows = book.active
    header = [cell.value for cell in rows[1]]
    for row in rows.iter_rows(min_row=2):
        if row[header.index("kind")].value == "answer":
            row[header.index("status")].value = "corrected"
            row[header.index("corrected_text")].value = "Igisubizo cyakosowe."
        else:
            row[header.index("status")].value = "approved"
    book.save(sheet)

    assert apply_decisions(sheet, output) == {"pending": 0, "approved": 1, "corrected": 1, "rejected": 0}
    records = list(master().values())
    examples = augmented_examples(records, output)
    assert [e["messages"][-1]["content"] for e in examples] == [pair["answer"], "Igisubizo cyakosowe."]
    assert all(e["id"] == pair["id"] and e["source"] == "gemini_reviewed" for e in examples)


def test_drafts_for_held_out_pairs_never_reach_training(tmp_path):
    held_out = next(r for r in master().values() if r["lang"] == "rw" and r["split"] == "test")
    path = tmp_path / "augmented.jsonl"
    path.write_text(json.dumps({"id": held_out["id"], "questions": ["Q?"], "answers": []}), encoding="utf-8")
    assert augmented_examples(list(master().values()), path) == []
