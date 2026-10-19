import json
from pathlib import Path

import pandas as pd
import pytest
from openpyxl import load_workbook

from nutriassist.qa_data.build_qa_dataset import SYSTEM_PROMPT, assign_splits, build_qa_dataset, original_records
from nutriassist.qa_data.review_sheet import apply_review, export_review_sheet, load_review


@pytest.fixture(scope="module")
def built(tmp_path_factory: pytest.TempPathFactory) -> Path:
    output = tmp_path_factory.mktemp("qa")
    build_qa_dataset(output_dir=output, review_path=output / "no_review.xlsx")
    return output


def _master(built: Path) -> list[dict]:
    return [json.loads(line) for line in (built / "qa_master.jsonl").open(encoding="utf-8")]


def test_every_original_pair_has_one_decision() -> None:
    records, dropped = original_records(
        Path("data/QA_data/malnutrition_dataset_final.jsonl"), Path("data/QA_data/curated/revisions.jsonl")
    )
    assert len(records) + len(dropped) == 135


def test_dropped_and_clinical_content_is_gone(built: Path) -> None:
    text = (built / "train.jsonl").read_text(encoding="utf-8").lower()
    for clinical in ("furosemide", "darrow", "5 ml/kg", "potassium permanganate", "cotrimoxazole"):
        assert clinical not in text


def test_groups_never_cross_splits(built: Path) -> None:
    splits_by_group: dict[str, set[str]] = {}
    for record in _master(built):
        splits_by_group.setdefault(record["group"], set()).add(record["split"])
    assert all(len(splits) == 1 for splits in splits_by_group.values())


def test_held_out_questions_do_not_appear_in_training(built: Path) -> None:
    train_questions = {
        json.loads(line)["messages"][1]["content"] for line in (built / "train.jsonl").open(encoding="utf-8")
    }
    for name in ("val.jsonl", "test.jsonl", "heldout_paraphrases.jsonl"):
        for line in (built / name).open(encoding="utf-8"):
            assert json.loads(line)["messages"][1]["content"] not in train_questions


def test_chat_format_has_system_user_assistant(built: Path) -> None:
    example = json.loads((built / "train.jsonl").open(encoding="utf-8").readline())
    assert [m["role"] for m in example["messages"]] == ["system", "user", "assistant"]
    assert example["messages"][0]["content"] == SYSTEM_PROMPT


def test_nutrient_answers_match_food_table(built: Path) -> None:
    nutrients = pd.read_csv("data/processed/food_nutrients.csv")
    protein = nutrients[(nutrients["food_id"] == "food_69e28b0cbcaf") & (nutrients["nutrient_name"] == "protein")]
    expected = f"{protein['amount_per_100g'].iloc[0]:.1f} g of protein"
    answers = [r["answer"] for r in _master(built) if r["id"] == "en-nut-69e28b0cbcaf"]
    assert answers and expected in answers[0]


def test_kinyarwanda_pairs_are_flagged_for_review(built: Path) -> None:
    rw = [r for r in _master(built) if r["lang"] == "rw"]
    assert rw and all(r["review"] == "pending_native_review" for r in rw)


def _rw(id_: str, question: str = "Q?", answer: str = "A.") -> dict:
    return {"id": id_, "group": id_, "lang": "rw", "category": "c", "question": question,
            "answer": answer, "paraphrases": [], "review": "pending_native_review"}


def _review(path: Path, decisions: dict[str, str], corrected_answer: str = "") -> None:
    export_review_sheet([_rw(i) for i in decisions], path)
    workbook = load_workbook(path)
    sheet = workbook.active
    header = [cell.value for cell in sheet[1]]
    for row in sheet.iter_rows(min_row=2):
        row[header.index("status")].value = decisions[row[0].value]
        row[header.index("rw_answer_corrected")].value = corrected_answer
    workbook.save(path)


def test_review_sheet_applies_reviewer_decisions(tmp_path: Path) -> None:
    sheet = tmp_path / "review.xlsx"
    _review(sheet, {"a": "approved", "b": "corrected", "c": "rejected", "d": "pending"}, corrected_answer="Igisubizo.")
    result = {r["id"]: r for r in apply_review([_rw(i) for i in "abcd"], sheet)}

    assert set(result) == {"a", "b", "d"}
    assert result["a"]["review"] == "approved"
    assert result["b"]["answer"] == "Igisubizo." and result["b"]["review"] == "corrected"
    assert result["d"]["review"] == "pending_native_review"
    assert set(r["id"] for r in apply_review([_rw(i) for i in "abcd"], sheet, reviewed_only=True)) == {"a", "b"}


def test_approval_is_reset_when_draft_changes(tmp_path: Path) -> None:
    sheet = tmp_path / "review.xlsx"
    _review(sheet, {"a": "approved"})
    changed = apply_review([_rw("a", answer="A different draft.")], sheet)
    assert changed[0]["review"] == "pending_native_review"


def test_reexport_keeps_existing_review(tmp_path: Path) -> None:
    sheet = tmp_path / "review.xlsx"
    _review(sheet, {"a": "approved", "b": "rejected"})
    export_review_sheet([_rw("a"), _rw("b"), _rw("new")], sheet)
    statuses = {k: v["status"] for k, v in load_review(sheet).items()}
    assert statuses == {"a": "approved", "b": "rejected", "new": "pending"}


def test_split_assignment_is_deterministic() -> None:
    records = [{"group": f"g{i}", "category": "c"} for i in range(40)]
    assert assign_splits(records) == assign_splits(records)
