"""The review of extra Kinyarwanda wordings: drafts in, reviewer decisions kept, approved wordings out."""

import json

import pytest
from openpyxl import load_workbook

from nutriassist.qa_data.wording_review import apply_decisions, export_sheet, load_drafts

MASTER = "data/QA_data/processed/qa_master.jsonl"


@pytest.fixture
def drafts(tmp_path):
    path = tmp_path / "drafts.jsonl"
    rows = [
        {"id": "rw-feed-13", "wordings": ["Icyayi ni kibi ku mwana?", "Icyayi  ni kibi ku mwana?", "Uruhinja rwanywa icyayi?"]},
        # The pair's own question is not a new wording.
        {"id": "rw-feed-06", "wordings": ["Nshobora guha igi umwana wanjye w'amezi 6?", "Igi ni ryiza ku mwana?"]},
    ]
    path.write_text("\n".join(json.dumps(r, ensure_ascii=False) for r in rows), encoding="utf-8")
    return path


def review(sheet, decisions):
    book = load_workbook(sheet)
    rows = book.active
    header = [cell.value for cell in rows[1]]
    for row in rows.iter_rows(min_row=2):
        wording = row[header.index("draft_wording")].value
        if wording in decisions:
            status, corrected = decisions[wording]
            row[header.index("status")].value = status
            row[header.index("corrected_wording")].value = corrected
    book.save(sheet)


def test_drafts_drop_duplicates_and_existing_wordings(drafts):
    assert load_drafts(drafts, MASTER) == [
        ("rw-feed-13", "Icyayi ni kibi ku mwana?"), ("rw-feed-13", "Uruhinja rwanywa icyayi?"),
        ("rw-feed-06", "Igi ni ryiza ku mwana?"),
    ]


def test_drafts_must_belong_to_a_kinyarwanda_pair(tmp_path):
    path = tmp_path / "bad.jsonl"
    path.write_text(json.dumps({"id": "en-feed-13", "wordings": ["Is tea bad?"]}), encoding="utf-8")
    with pytest.raises(ValueError, match="not a Kinyarwanda pair"):
        load_drafts(path, MASTER)


def test_only_approved_and_corrected_wordings_reach_the_app(drafts, tmp_path):
    sheet, output = tmp_path / "review.xlsx", tmp_path / "approved.jsonl"
    export_sheet(drafts, sheet, MASTER)
    review(sheet, {"Icyayi ni kibi ku mwana?": ("approved", None),
                   "Uruhinja rwanywa icyayi?": ("rejected", None),
                   "Igi ni ryiza ku mwana?": ("corrected", "Igi ni ryiza ku mwana muto?")})
    counts = apply_decisions(sheet, output)
    assert counts == {"pending": 0, "approved": 1, "corrected": 1, "rejected": 1}
    rows = [json.loads(line) for line in output.read_text(encoding="utf-8").splitlines()]
    assert rows == [{"id": "rw-feed-06", "wordings": ["Igi ni ryiza ku mwana muto?"]},
                    {"id": "rw-feed-13", "wordings": ["Icyayi ni kibi ku mwana?"]}]


def test_re_export_keeps_decisions(drafts, tmp_path):
    sheet = tmp_path / "review.xlsx"
    export_sheet(drafts, sheet, MASTER)
    review(sheet, {"Icyayi ni kibi ku mwana?": ("approved", None)})
    export_sheet(drafts, sheet, MASTER)
    statuses = {row[4]: row[5] for row in load_workbook(sheet).active.iter_rows(min_row=2, values_only=True)}
    assert statuses["Icyayi ni kibi ku mwana?"] == "approved" and statuses["Igi ni ryiza ku mwana?"] == "pending"


def test_drafted_wordings_file_is_valid():
    drafts = load_drafts()
    assert len(drafts) >= 100 and all(pair_id.startswith("rw-") for pair_id, _ in drafts)
