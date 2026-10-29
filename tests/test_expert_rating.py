import pandas as pd
import pytest
from openpyxl import load_workbook

from nutriassist.evaluation import expert_rating as er


@pytest.fixture
def predictions(tmp_path):
    rows = []
    for i in range(8):
        for lang in ("en", "rw"):
            rows.append({"id": f"{lang}-{i}", "lang": lang, "category": f"cat{i % 3}", "question": f"Q{i} {lang}",
                         "reference": "ref", "base_answer": f"base {i}", "fine_tuned_answer": f"tuned {i}"})
    path = tmp_path / "qualitative_test.csv"
    pd.DataFrame(rows).to_csv(path, index=False)
    return path


def fill(sheet, key, score_of):
    """Score every row of the sheet like an expert would: score_of(model) -> 1..5."""
    models = pd.read_csv(key).set_index("item")["model"]
    workbook = load_workbook(sheet)
    ratings = workbook["Ratings"]
    for row in range(2, ratings.max_row + 1):
        score = score_of(models[ratings.cell(row, 1).value])
        for col in range(4, 7):
            ratings.cell(row, col).value = score
    workbook.save(sheet)


def test_sheet_is_blinded_and_covers_every_category(predictions, tmp_path):
    sheet, key = tmp_path / "rating.xlsx", tmp_path / "key.csv"
    rows = er.build(predictions, sheet, key)

    header = [c.value for c in load_workbook(sheet)["Ratings"][1]]
    assert "model" not in header and header[3:6] == list(er.CRITERIA)   # the expert never sees which model
    assert len(rows) == 16 and set(pd.read_csv(key)["model"]) == {"base", "fine_tuned"}
    assert all(q.endswith("en") for q in rows["question"])               # English answers only
    assert pd.read_csv(predictions).set_index("id").loc[rows["id"].unique(), "category"].nunique() == 3


def test_analysis_compares_the_models(predictions, tmp_path):
    sheet, key = tmp_path / "rating.xlsx", tmp_path / "key.csv"
    er.build(predictions, sheet, key)
    fill(sheet, key, lambda model: 5 if model == "fine_tuned" else 2)

    ratings = er.read_ratings(sheet, key)
    table = er.summary(ratings)
    assert table.loc["fine_tuned", "correctness"] == "5.00 ± 0.00"
    assert table.loc["base", "rated_4_or_5_%"] == 0
    assert all(p < 0.05 for p in er.paired_tests(ratings).values())
    assert er.agreement(ratings, ratings)["safety"] == pytest.approx(1.0)


def test_unrated_rows_are_left_out(predictions, tmp_path):
    sheet, key = tmp_path / "rating.xlsx", tmp_path / "key.csv"
    er.build(predictions, sheet, key)
    assert er.read_ratings(sheet, key).empty
