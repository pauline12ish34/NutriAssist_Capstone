"""Gemini drafts of extra Kinyarwanda training text, used only after a native speaker approves them.

The fine-tuned model writes poor Kinyarwanda mostly because it has seen little of it (181 training
pairs). For each reviewed Kinyarwanda pair in the training split, Gemini drafts new ways a caregiver
might ask the question and one reworded answer that keeps the same facts:

    python -m nutriassist.qa_data.gemini_drafts draft    # calls Gemini -> curated/gemini_drafts_rw.jsonl
    python -m nutriassist.qa_data.gemini_drafts export   # writes the review sheet
    python -m nutriassist.qa_data.gemini_drafts apply    # approved drafts -> processed/augmented_rw.jsonl

``draft`` can be stopped and rerun: pairs already drafted are skipped. In the sheet the reviewer sets
``status`` to ``approved``, ``corrected`` (and fills ``corrected_text``) or ``rejected``; re-exporting
keeps every decision already made. The ``check`` column flags drafts whose numbers differ from the
reviewed answer or that use words the reviewer has ruled out.

The dataset builder adds approved drafts to the training split only (each new question with the
reviewed answer, each reworded answer with the reviewed question), so validation and test data are
unchanged and the scores stay comparable with earlier runs.
Approved question wordings also reach the app's matching (``nutriassist.knowledge``), so it
recognises more ways of asking a reviewed question.

Drafting needs ``pip install google-genai`` and an API key from Google AI Studio, either in the
environment (``GEMINI_API_KEY`` or ``GOOGLE_API_KEY``) or in ``.env`` (which git ignores).
"""

from __future__ import annotations

import argparse
import json
import os
import re
import time
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from ..terms import FOODS, GLOSSARY, ILLNESSES

QA_DIR = Path("data/QA_data")
MASTER = QA_DIR / "processed" / "qa_master.jsonl"
DRAFTS = QA_DIR / "curated" / "gemini_drafts_rw.jsonl"
SHEET = QA_DIR / "review" / "gemini_drafts_review.xlsx"
AUGMENTED = QA_DIR / "processed" / "augmented_rw.jsonl"
ENV_FILE = Path(".env")
KEY_NAMES = ("GEMINI_API_KEY", "GOOGLE_API_KEY", "GOOGLE_APIKEY")

DEFAULT_MODEL = "gemini-3.5-flash"   # 3.8 flash was overloaded (503/504) on the free tier; pro has no free quota
QUESTIONS_PER_PAIR = 3
BATCH_SIZE = 5           # pairs per request: fewer requests stay inside the free tier's daily limit
DELAY_SECONDS = 7.0      # between requests: the free tier allows about 10 a minute
MAX_RETRIES = 5
REQUEST_TIMEOUT_MS = 120_000

# The reviewer replaced gusekura (pound) with gusya (grind, mash) in every answer.
BANNED_WORDS = ("gusekura", "isekuye", "bisekuye", "usekura")
STATUSES = ("pending", "approved", "corrected", "rejected")
KINDS = ("question", "answer")
COLUMNS = [("id", 14), ("category", 16), ("kind", 10), ("english_question", 36), ("rw_question", 36),
           ("rw_answer", 50), ("draft", 50), ("check", 22), ("status", 12), ("corrected_text", 50),
           ("notes", 26)]
REVIEWER_COLUMNS = ("status", "corrected_text", "notes")
_NUMBER = re.compile(r"\d+(?:[.,]\d+)?")


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _normalise(text: str) -> str:
    return " ".join(text.lower().split())


def _read_jsonl(path: str | Path) -> list[dict]:
    with Path(path).open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def api_key(env_path: str | Path = ENV_FILE) -> str:
    """The Gemini key from the environment, else from ``.env`` (names are matched case-insensitively)."""
    for name in KEY_NAMES:
        if os.environ.get(name):
            return os.environ[name]
    env_path = Path(env_path)
    if env_path.exists():
        for line in env_path.read_text(encoding="utf-8").splitlines():
            name, _, value = line.partition("=")
            if name.strip().upper() in KEY_NAMES and value.strip():
                return value.strip().strip("'\"")
    raise RuntimeError(f"No Gemini API key: add GEMINI_API_KEY=... to {env_path} or the environment")


def source_pairs(master_path: str | Path = MASTER) -> list[dict]:
    """Reviewed Kinyarwanda training pairs, each with its English source."""
    master = _read_jsonl(master_path)
    english = {r["group"]: r for r in master if r["lang"] == "en"}
    return [
        {**r, "english_question": english.get(r["group"], {}).get("question", ""),
         "english_answer": english.get(r["group"], {}).get("answer", "")}
        for r in master
        if r["lang"] == "rw" and r["split"] == "train" and r["review"] in ("approved", "corrected")
    ]


def _glossary() -> str:
    terms = [t for t in [*GLOSSARY, *ILLNESSES, *FOODS] if t.translate]
    return "\n".join(f"- {t.english[0]} = {t.kinyarwanda[0]}" for t in terms)


def build_prompt(pairs: list[dict], questions_per_pair: int = QUESTIONS_PER_PAIR) -> str:
    items = [{"id": p["id"], "english_question": p["english_question"], "english_answer": p["english_answer"],
              "kinyarwanda_question": p["question"], "kinyarwanda_other_wordings": p["paraphrases"],
              "kinyarwanda_answer": p["answer"]} for p in pairs]
    return f"""You help build training data for a child-nutrition assistant for caregivers in Rwanda.
Each item below is a question and answer that a native Kinyarwanda speaker has already checked.

For each item write, in natural everyday Kinyarwanda as spoken in Rwanda (not Kirundi, not English):
1. "questions": {questions_per_pair} new ways a parent or grandmother might ask the same question. Vary them:
   short and long, informal, different words. Each must still be answered by the same answer.
   Do not repeat the existing wordings.
2. "answer": the Kinyarwanda answer reworded. Keep every fact, number, age and unit exactly. Keep any
   advice to go to the health centre (ikigo nderabuzima). Add nothing new and leave nothing out.

Use these words:
{_glossary()}
- to mash or grind food = gusya (never gusekura)

Reply with only a JSON array, one object per item: {{"id": "...", "questions": ["..."], "answer": "..."}}

Items:
{json.dumps(items, ensure_ascii=False, indent=1)}"""


def parse_response(text: str, ids: set[str]) -> dict[str, dict]:
    """Drafts by pair id from Gemini's JSON reply, ignoring unknown ids and malformed entries."""
    text = text.strip()
    if text.startswith("```"):
        text = text.strip("`").removeprefix("json").strip()
    drafts = {}
    for item in json.loads(text):
        if not isinstance(item, dict) or item.get("id") not in ids:
            continue
        questions = [" ".join(q.split()) for q in item.get("questions", []) if isinstance(q, str) and q.strip()]
        answer = item.get("answer")
        drafts[item["id"]] = {"questions": questions,
                              "answers": [" ".join(answer.split())] if isinstance(answer, str) and answer.strip() else []}
    return drafts


def _generate(client, model: str, prompt: str) -> str:
    import httpx
    from google.genai import errors, types

    config = types.GenerateContentConfig(temperature=0.8, response_mime_type="application/json")
    for attempt in range(MAX_RETRIES):
        try:
            return client.models.generate_content(model=model, contents=prompt, config=config).text
        except (errors.APIError, httpx.TransportError) as error:
            # 429: over the rate limit; 5xx: Google is busy; a dropped connection: the network. All pass if we wait.
            code = getattr(error, "code", None)
            if (code is not None and code != 429 and code < 500) or attempt == MAX_RETRIES - 1:
                raise
            wait = 30 * 2 ** attempt
            print(f"  {f'Gemini error {code}' if code else 'connection failed'}; waiting {wait}s")
            time.sleep(wait)
    raise AssertionError("unreachable")


def draft(master_path: str | Path = MASTER, drafts_path: str | Path = DRAFTS, model: str = DEFAULT_MODEL,
          batch_size: int = BATCH_SIZE, delay: float = DELAY_SECONDS, limit: int | None = None) -> int:
    """Ask Gemini for drafts of the pairs not drafted yet. Returns the number of pairs drafted."""
    from google import genai

    drafts_path = Path(drafts_path)
    done = {row["id"] for row in _read_jsonl(drafts_path)} if drafts_path.exists() else set()
    todo = [p for p in source_pairs(master_path) if p["id"] not in done][:limit]
    # Without a timeout a dropped connection can hang a request for good.
    client = genai.Client(api_key=api_key(), http_options=genai.types.HttpOptions(timeout=REQUEST_TIMEOUT_MS))
    drafts_path.parent.mkdir(parents=True, exist_ok=True)

    drafted = 0
    for start in range(0, len(todo), batch_size):
        batch = todo[start:start + batch_size]
        print(f"Drafting {start + 1}-{start + len(batch)} of {len(todo)}")
        reply = _generate(client, model, build_prompt(batch))
        try:
            results = parse_response(reply, {p["id"] for p in batch})
        except json.JSONDecodeError:
            print("  Gemini's reply was not valid JSON; these pairs will be retried next run")
            results = {}
        with drafts_path.open("a", encoding="utf-8") as handle:
            for pair_id, result in results.items():
                handle.write(json.dumps({"id": pair_id, "model": model, **result}, ensure_ascii=False) + "\n")
        drafted += len(results)
        if start + batch_size < len(todo):
            time.sleep(delay)
    return drafted


def checks(kind: str, text: str, pair: dict) -> str:
    """Warnings for the reviewer: changed numbers or ruled-out words."""
    warnings = []
    if kind == "answer" and sorted(_NUMBER.findall(text)) != sorted(_NUMBER.findall(pair["answer"])):
        warnings.append("numbers differ")
    if any(word in text.lower() for word in BANNED_WORDS):
        warnings.append("uses gusekura")
    return "; ".join(warnings)


def load_drafts(drafts_path: str | Path = DRAFTS, master_path: str | Path = MASTER) -> list[tuple[str, str, str]]:
    """(pair id, kind, text) in file order, without duplicates or questions the dataset already has."""
    master = {row["id"]: row for row in _read_jsonl(master_path)}
    known = {_normalise(q) for row in master.values() for q in [row["question"], *row["paraphrases"]]}
    answers = {_normalise(row["answer"]) for row in master.values()}
    drafts, seen = [], set()
    for row in _read_jsonl(drafts_path):
        pair = master.get(row["id"])
        if pair is None or pair["lang"] != "rw":
            raise ValueError(f"{row['id']} in {drafts_path} is not a Kinyarwanda pair in {master_path}")
        for kind, existing in (("question", known), ("answer", answers)):
            for text in row.get(f"{kind}s", []):
                text = " ".join(text.split())
                key = (row["id"], kind, _normalise(text))
                if text and key not in seen and _normalise(text) not in existing:
                    seen.add(key)
                    drafts.append((row["id"], kind, text))
    return drafts


def load_decisions(path: str | Path = SHEET) -> dict[tuple[str, str, str], dict[str, str]]:
    """Reviewer decisions keyed by (pair id, kind, draft text)."""
    path = Path(path)
    if not path.exists():
        return {}
    rows = load_workbook(path, read_only=True).active.iter_rows(values_only=True)
    header = [_clean(cell) for cell in next(rows)]
    decisions = {}
    for row in rows:
        entry = {name: _clean(value) for name, value in zip(header, row)}
        if not entry.get("id"):
            continue
        entry["status"] = entry.get("status", "").lower() or "pending"
        if entry["status"] not in STATUSES:
            raise ValueError(f"Unknown status {entry['status']!r} for {entry['id']} in {path}")
        decisions[(entry["id"], entry["kind"], entry["draft"])] = entry
    return decisions


def export_sheet(drafts_path: str | Path = DRAFTS, sheet_path: str | Path = SHEET,
                 master_path: str | Path = MASTER) -> Path:
    sheet_path = Path(sheet_path)
    master = {row["id"]: row for row in _read_jsonl(master_path)}
    english = {row["group"]: row["question"] for row in master.values() if row["lang"] == "en"}
    decisions = load_decisions(sheet_path)

    book = Workbook()
    sheet = book.active
    sheet.title = "gemini_drafts"
    sheet.append([name for name, _ in COLUMNS])
    for pair_id, kind, text in load_drafts(drafts_path, master_path):
        pair = master[pair_id]
        previous = decisions.get((pair_id, kind, text), {})
        sheet.append([pair_id, pair["category"], kind, english.get(pair["group"], ""), pair["question"],
                      pair["answer"], text, checks(kind, text, pair), previous.get("status") or "pending",
                      previous.get("corrected_text", ""), previous.get("notes", "")])

    header_fill, reviewer_fill = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FFF2CC")
    for column, (name, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, column)
        sheet.column_dimensions[cell.column_letter].width = width
        cell.font = Font(bold=True)
        cell.fill = reviewer_fill if name in REVIEWER_COLUMNS else header_fill
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    status_letter = sheet.cell(1, [n for n, _ in COLUMNS].index("status") + 1).column_letter
    validation = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=False)
    validation.add(f"{status_letter}2:{status_letter}{max(sheet.max_row, 2)}")
    sheet.add_data_validation(validation)
    sheet.freeze_panes = "G2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(sheet_path)
    return sheet_path


def apply_decisions(sheet_path: str | Path = SHEET, output: str | Path = AUGMENTED) -> dict[str, int]:
    """Write approved and corrected drafts for the dataset builder. Returns counts per status."""
    output = Path(output)
    approved: dict[str, dict[str, list[str]]] = {}
    counts = dict.fromkeys(STATUSES, 0)
    for (pair_id, kind, text), entry in load_decisions(sheet_path).items():
        counts[entry["status"]] += 1
        if entry["status"] == "corrected" and entry["corrected_text"]:
            text = entry["corrected_text"]
        elif entry["status"] != "approved":
            continue
        approved.setdefault(pair_id, {"questions": [], "answers": []})[f"{kind}s"].append(text)
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8") as handle:
        for pair_id, texts in sorted(approved.items()):
            handle.write(json.dumps({"id": pair_id, **texts}, ensure_ascii=False) + "\n")
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("action", choices=["draft", "export", "apply"])
    parser.add_argument("--model", default=DEFAULT_MODEL, help=f"Gemini model (default {DEFAULT_MODEL})")
    parser.add_argument("--batch-size", type=int, default=BATCH_SIZE, help="pairs per Gemini request")
    parser.add_argument("--delay", type=float, default=DELAY_SECONDS, help="seconds between requests")
    parser.add_argument("--limit", type=int, help="draft at most this many pairs (for a trial run)")
    args = parser.parse_args()
    if args.action == "draft":
        print(f"Drafted {draft(model=args.model, batch_size=args.batch_size, delay=args.delay, limit=args.limit)} "
              f"pairs into {DRAFTS}")
    elif args.action == "export":
        print(f"Wrote {export_sheet()} ({len(load_drafts())} drafts)")
    else:
        counts = apply_decisions()
        print(f"Wrote {AUGMENTED}: " + ", ".join(f"{n} {status}" for status, n in counts.items()))


if __name__ == "__main__":
    main()
