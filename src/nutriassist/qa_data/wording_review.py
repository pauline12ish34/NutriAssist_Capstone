"""Native-speaker review of extra Kinyarwanda question wordings, which help the app recognise questions.

The app answers Kinyarwanda only from reviewed answers, so it must recognise a question however a
caregiver words it. Extra wordings for the reviewed pairs are drafted in
``data/QA_data/curated/extra_wordings_rw.jsonl`` (one row per pair: ``{"id": ..., "wordings": [...]}``)
and reviewed one by one:

    python -m nutriassist.qa_data.wording_review export   # writes the review sheet
    python -m nutriassist.qa_data.wording_review apply    # approved wordings -> processed/extra_wordings.jsonl

In the sheet the reviewer sets ``status`` to ``approved``, ``corrected`` (and fills ``corrected_wording``)
or ``rejected``. Re-exporting keeps every decision already made. Only approved and corrected wordings
reach the app (``nutriassist.knowledge``). They are not used for training, so adding them does not
change the training data or the reported evaluation.
"""

from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

QA_DIR = Path("data/QA_data")
DRAFTS = QA_DIR / "curated" / "extra_wordings_rw.jsonl"
SHEET = QA_DIR / "review" / "kinyarwanda_wordings_review.xlsx"
APPROVED = QA_DIR / "processed" / "extra_wordings.jsonl"
MASTER = QA_DIR / "processed" / "qa_master.jsonl"
STATUSES = ("pending", "approved", "corrected", "rejected")
COLUMNS = [("id", 14), ("category", 16), ("english_question", 38), ("rw_question", 38), ("draft_wording", 42),
           ("status", 12), ("corrected_wording", 42), ("notes", 28)]
REVIEWER_COLUMNS = ("status", "corrected_wording", "notes")


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def _read_jsonl(path: str | Path) -> list[dict]:
    with Path(path).open(encoding="utf-8") as handle:
        return [json.loads(line) for line in handle if line.strip()]


def load_drafts(path: str | Path = DRAFTS, master_path: str | Path = MASTER) -> list[tuple[str, str]]:
    """(pair id, draft wording) in file order, without duplicates or wordings the pair already has."""
    master = {row["id"]: row for row in _read_jsonl(master_path)}
    drafts, seen = [], set()
    for row in _read_jsonl(path):
        pair = master.get(row["id"])
        if pair is None or pair["lang"] != "rw":
            raise ValueError(f"{row['id']} in {path} is not a Kinyarwanda pair in {master_path}")
        known = {pair["question"].lower(), *(p.lower() for p in pair["paraphrases"])}
        for wording in row["wordings"]:
            wording = " ".join(wording.split())
            key = (row["id"], wording.lower())
            if wording and key not in seen and wording.lower() not in known:
                seen.add(key)
                drafts.append((row["id"], wording))
    return drafts


def load_decisions(path: str | Path = SHEET) -> dict[tuple[str, str], dict[str, str]]:
    """Reviewer decisions keyed by (pair id, draft wording)."""
    path = Path(path)
    if not path.exists():
        return {}
    rows = load_workbook(path, read_only=True).active.iter_rows(values_only=True)
    header = [_clean(cell) for cell in next(rows)]
    decisions = {}
    for row in rows:
        entry = {name: _clean(value) for name, value in zip(header, row)}
        if not entry.get("id"):
            continue
        entry["status"] = entry.get("status", "").lower() or "pending"
        if entry["status"] not in STATUSES:
            raise ValueError(f"Unknown status {entry['status']!r} for {entry['id']} in {path}")
        decisions[(entry["id"], entry["draft_wording"])] = entry
    return decisions


def export_sheet(drafts_path: str | Path = DRAFTS, sheet_path: str | Path = SHEET,
                 master_path: str | Path = MASTER) -> Path:
    sheet_path = Path(sheet_path)
    master = {row["id"]: row for row in _read_jsonl(master_path)}
    english = {row["group"]: row["question"] for row in master.values() if row["lang"] == "en"}
    decisions = load_decisions(sheet_path)

    book = Workbook()
    sheet = book.active
    sheet.title = "wordings"
    sheet.append([name for name, _ in COLUMNS])
    for pair_id, wording in load_drafts(drafts_path, master_path):
        pair = master[pair_id]
        previous = decisions.get((pair_id, wording), {})
        sheet.append([pair_id, pair["category"], english.get(pair["group"], ""), pair["question"], wording,
                      previous.get("status") or "pending", previous.get("corrected_wording", ""),
                      previous.get("notes", "")])

    header_fill, reviewer_fill = PatternFill("solid", fgColor="DDEBF7"), PatternFill("solid", fgColor="FFF2CC")
    for column, (name, width) in enumerate(COLUMNS, start=1):
        cell = sheet.cell(1, column)
        sheet.column_dimensions[cell.column_letter].width = width
        cell.font = Font(bold=True)
        cell.fill = reviewer_fill if name in REVIEWER_COLUMNS else header_fill
    for cells in sheet.iter_rows(min_row=2):
        for cell in cells:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    status_letter = sheet.cell(1, [n for n, _ in COLUMNS].index("status") + 1).column_letter
    validation = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=False)
    validation.add(f"{status_letter}2:{status_letter}{max(sheet.max_row, 2)}")
    sheet.add_data_validation(validation)
    sheet.freeze_panes = "E2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet_path.parent.mkdir(parents=True, exist_ok=True)
    book.save(sheet_path)
    return sheet_path


def apply_decisions(sheet_path: str | Path = SHEET, output: str | Path = APPROVED) -> dict[str, int]:
    """Write approved and corrected wordings for the app. Returns counts per status."""
    output = Path(output)
    approved: dict[str, list[str]] = {}
    counts = dict.fromkeys(STATUSES, 0)
    for (pair_id, draft), entry in load_decisions(sheet_path).items():
        counts[entry["status"]] += 1
        if entry["status"] == "approved":
            approved.setdefault(pair_id, []).append(draft)
        elif entry["status"] == "corrected" and entry["corrected_wording"]:
            approved.setdefault(pair_id, []).append(entry["corrected_wording"])
    output.parent.mkdir(parents=True, exist_ok=True)
    with output.open("w", encoding="utf-8") as handle:
        for pair_id, wordings in sorted(approved.items()):
            handle.write(json.dumps({"id": pair_id, "wordings": wordings}, ensure_ascii=False) + "\n")
    return counts


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("action", choices=["export", "apply"])
    args = parser.parse_args()
    if args.action == "export":
        print(f"Wrote {export_sheet()} ({len(load_drafts())} wordings)")
    else:
        counts = apply_decisions()
        print(f"Wrote {APPROVED}: " + ", ".join(f"{n} {status}" for status, n in counts.items()))


if __name__ == "__main__":
    main()
