"""Native-speaker review of the Kinyarwanda QA pairs through an Excel sheet.

``python -m nutriassist.qa_data.review_sheet`` writes ``data/QA_data/review/kinyarwanda_review.xlsx``.
The reviewer sets ``status`` on each row and, when needed, fills the ``corrected``
columns. Re-exporting keeps every review already made. The dataset builder reads
the sheet back with :func:`apply_review`.

An approval only holds for the draft text the reviewer saw: if a draft changes
after review, the row goes back to ``pending``.
"""

from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

REVIEW_SHEET = Path("data/QA_data/review/kinyarwanda_review.xlsx")
STATUSES = ("pending", "approved", "corrected", "rejected")
PARAPHRASE_SEPARATOR = "\n"

COLUMNS = [
    ("id", 16),
    ("category", 16),
    ("english_question", 40),
    ("english_answer", 60),
    ("rw_question_draft", 40),
    ("rw_answer_draft", 60),
    ("rw_paraphrases_draft", 35),
    ("status", 12),
    ("rw_question_corrected", 40),
    ("rw_answer_corrected", 60),
    ("rw_paraphrases_corrected", 35),
    ("reviewer_notes", 30),
]
REVIEWER_COLUMNS = ("status", "rw_question_corrected", "rw_answer_corrected", "rw_paraphrases_corrected", "reviewer_notes")


def _clean(value: object) -> str:
    return "" if value is None else str(value).strip()


def load_review(path: str | Path = REVIEW_SHEET) -> dict[str, dict[str, str]]:
    path = Path(path)
    if not path.exists():
        return {}
    sheet = load_workbook(path, read_only=True).active
    rows = sheet.iter_rows(values_only=True)
    header = [_clean(cell) for cell in next(rows)]
    review = {}
    for row in rows:
        entry = {name: _clean(value) for name, value in zip(header, row)}
        if entry.get("id"):
            entry["status"] = entry.get("status", "").lower() or "pending"
            if entry["status"] not in STATUSES:
                raise ValueError(f"Unknown status {entry['status']!r} for {entry['id']} in {path}")
            review[entry["id"]] = entry
    return review


def _draft_matches(record: dict, entry: dict[str, str]) -> bool:
    return (
        entry.get("rw_question_draft") == record["question"]
        and entry.get("rw_answer_draft") == record["answer"]
        and entry.get("rw_paraphrases_draft", "") == PARAPHRASE_SEPARATOR.join(record["paraphrases"])
    )


def apply_review(records: list[dict], path: str | Path = REVIEW_SHEET, reviewed_only: bool = False) -> list[dict]:
    """Apply reviewer decisions to Kinyarwanda records. English records pass through unchanged."""
    review = load_review(path)
    kept = []
    for record in records:
        if record["lang"] != "rw":
            kept.append(record)
            continue
        entry = review.get(record["id"])
        status = entry["status"] if entry and _draft_matches(record, entry) else "pending"
        if status == "rejected" or (reviewed_only and status == "pending"):
            continue
        if status == "corrected":
            record["question"] = entry["rw_question_corrected"] or record["question"]
            record["answer"] = entry["rw_answer_corrected"] or record["answer"]
            if entry["rw_paraphrases_corrected"]:
                record["paraphrases"] = [p.strip() for p in entry["rw_paraphrases_corrected"].split(PARAPHRASE_SEPARATOR) if p.strip()]
        record["review"] = "pending_native_review" if status == "pending" else status
        kept.append(record)
    return kept


def export_review_sheet(records: list[dict], path: str | Path = REVIEW_SHEET) -> Path:
    """Write the review sheet, keeping decisions already made for unchanged drafts."""
    path = Path(path)
    existing = load_review(path)
    english = {r["group"]: r for r in records if r["lang"] == "en"}

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "kinyarwanda_review"
    sheet.append([name for name, _ in COLUMNS])
    for record in (r for r in records if r["lang"] == "rw"):
        source = english.get(record["group"], {})
        entry = existing.get(record["id"])
        previous = entry if entry and _draft_matches(record, entry) else {}
        values = {
            "id": record["id"],
            "category": record["category"],
            "english_question": source.get("question", ""),
            "english_answer": source.get("answer", ""),
            "rw_question_draft": record["question"],
            "rw_answer_draft": record["answer"],
            "rw_paraphrases_draft": PARAPHRASE_SEPARATOR.join(record["paraphrases"]),
            **{column: previous.get(column, "") for column in REVIEWER_COLUMNS},
        }
        values["status"] = values["status"] or "pending"
        sheet.append([values[name] for name, _ in COLUMNS])

    header_fill = PatternFill("solid", fgColor="DDEBF7")
    reviewer_fill = PatternFill("solid", fgColor="FFF2CC")
    for index, (name, width) in enumerate(COLUMNS, start=1):
        letter = sheet.cell(row=1, column=index).column_letter
        sheet.column_dimensions[letter].width = width
        sheet.cell(row=1, column=index).font = Font(bold=True)
        sheet.cell(row=1, column=index).fill = reviewer_fill if name in REVIEWER_COLUMNS else header_fill
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True, vertical="top")
    status_letter = sheet.cell(row=1, column=[n for n, _ in COLUMNS].index("status") + 1).column_letter
    validation = DataValidation(type="list", formula1=f'"{",".join(STATUSES)}"', allow_blank=False)
    validation.add(f"{status_letter}2:{status_letter}{sheet.max_row}")
    sheet.add_data_validation(validation)
    sheet.freeze_panes = "C2"
    sheet.auto_filter.ref = sheet.dimensions

    path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(path)
    return path


if __name__ == "__main__":
    from .build_qa_dataset import collect_records

    records, _ = collect_records()
    print(f"Wrote {export_review_sheet(records)}")
